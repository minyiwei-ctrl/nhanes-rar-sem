"""Verified reporting of the bounded cycle and cognition-selection audit models
-> Supplementary Tables S22 (cycle sensitivity) and S23 (cognitive
availability and complete-case analysis).

    python freeze_sensitivity_report.py

Imported unchanged from the F1 results freeze (2026-10-19); run from stage
19 after the cycle_audit, cycle_contrasts and selection_audit models have
been run with 9_mplus_models/mplus_run.py (every output is read through
mplus_output_qc.read_verified, so a .run.json record with status PASS and
matching hashes is required). Also writes the full numerical audit to
result/qc/freeze/ (cycle_comparisons.csv, cycle_contrasts.csv,
cycle_pooled_contrasts.csv, selection_comparisons.csv,
sensitivity_verified_outputs.csv, sensitivity_summary.json).
"""
from pathlib import Path
from functools import lru_cache
import sys,json,math,re
import numpy as np
import pandas as pd
from openpyxl.styles import Font,PatternFill,Alignment
P=Path(__file__).resolve().parents[1];M=P/'scripts/9_mplus_models';Q=P/'result/qc/freeze'
sys.path.insert(0,str(P/'scripts'))
from mplus_parse import parse_estimates,parse_effects
from mplus_output_qc import read_verified,wald_result
from mplus_export import fmt_ci,fmt_p
QC=[]
@lru_cache(None)
def read(folder,stem):
    t,q=read_verified(M/folder/(stem+'.inp'));QC.append({'folder':folder,'model':stem,**q})
    return pd.DataFrame(parse_estimates(t,stem)),pd.DataFrame(parse_effects(t,stem)),q,t

def quantities(e,f,stem,sol):
    x=next(x for x in ['RA','IS','IV'] if re.search(r'_'+x.lower()+r'(?:_|$)',stem))
    if '_m_' in '_'+stem:
        y='COG' if '_cog_' in stem else 'PHQ9'
        specs=[('a1','BA_ADV ON',x),('a2','SII_LOG ON',x),('b1',y+' ON','BA_ADV'),('b2',y+' ON','SII_LOG')]
        out={label:e[(e.solution==sol)&(e.section==sec)&(e.parameter==pa)].iloc[0] for label,sec,pa in specs}
        for label,ef,via in [('direct','direct',None),('indirect_KDM','indirect','BA_ADV'),('indirect_SII','indirect','SII_LOG'),('total','total',None)]:
            a=f[(f.solution==sol)&(f.effect==ef)];a=a[a.via==via] if via else a[a.via.isna()]
            assert len(a)==1;out[label]=a.iloc[0]
    else:
        y='SII_LOG' if '_sii_' in stem else 'BA_ADV'
        a=e[(e.solution==sol)&(e.section==y+' ON')&(e.parameter==x+'_AGE')];assert len(a)==1
        out={'age_interaction':a.iloc[0]}
    return out

def compare(folder,stem,base,variant,sol):
    e,f,q,t=read(folder,stem);eb,fb,qb,tb=read('',base)
    a=quantities(e,f,stem,sol);b=quantities(eb,fb,base,sol);rows=[]
    for name,r in a.items():
        z={'model':base,'variant':variant,'solution':sol,'quantity':name,'n_base':qb['n'],'n_variant':q['n']}
        for prefix,r0 in [('base',b[name]),('variant',r)]:
            for c in ['estimate','se','ci_lo','ci_hi','p']:z[prefix+'_'+c]=float(r0[c])
        z['delta']=z['variant_estimate']-z['base_estimate'];rows.append(z)
    return rows

def book(path,sheets,notes):
    # Compact publication sheets retain links to the full numerical audit sheets.
    if 'A cycle adjustment' in sheets:
        a=sheets['A cycle adjustment']; a=a[(a.solution=='STDYX')&a.quantity.isin(['a1','a2','indirect_KDM','indirect_SII'])]
        compact=[]
        for _,r in a.iterrows():
            compact.append({'Model':r.model,'Effect (STDYX)':r.quantity,'n':int(r.n_base),'Baseline beta (95% CI)':fmt_ci(r.base_estimate,r.base_ci_lo,r.base_ci_hi),'Cycle-adjusted beta (95% CI)':fmt_ci(r.variant_estimate,r.variant_ci_lo,r.variant_ci_hi),'P adjusted':fmt_p(r.variant_p)})
        sheets['Word A primary effects']=pd.DataFrame(compact)
        compact=[]
        for _,r in sheets['A cycle adjustment'].query("quantity=='age_interaction'").iterrows():
            compact.append({'Model':r.model,'n':int(r.n_base),'Baseline b (95% CI)':fmt_ci(r.base_estimate,r.base_ci_lo,r.base_ci_hi),'Cycle-adjusted b (95% CI)':fmt_ci(r.variant_estimate,r.variant_ci_lo,r.variant_ci_hi),'P adjusted':fmt_p(r.variant_p)})
        sheets['Word B age interactions']=pd.DataFrame(compact)
        b=sheets['B functional form'].dropna(subset=['p']).copy();b['P']=b.p.map(lambda x:'<0.0001' if x<.0001 else f'{x:.4f}')
        sheets['Word C functional tests']=b[['model','variant','n','wald','df','P']]
        compact=[]
        for _,r in sheets['C pooled cycle contrasts'].iterrows():
            compact.append({'Age':r.age,'Outcome':r.outcome,'Metric':r.exposure,'n':int(r.n),'2011-12 b (95% CI)':fmt_ci(r.G_estimate,r.G_ci_lo,r.G_ci_hi),'2013-14 b (95% CI)':fmt_ci(r.H_estimate,r.H_ci_lo,r.H_ci_hi),'Difference b (95% CI)':fmt_ci(r.difference_estimate,r.difference_ci_lo,r.difference_ci_hi),'Wald P':f'{r.wald_p:.4f}' if r.wald_p>=.0001 else '<0.0001'})
        sheets['Word D cycle contrasts']=pd.DataFrame(compact)
    if 'C complete case' in sheets:
        compact=[]
        for _,r in sheets['C complete case'].query("solution=='STDYX'").iterrows():
            compact.append({'Model':r.model,'Effect (STDYX)':r.quantity,'n FIML':int(r.n_base),'n complete':int(r.n_variant),'FIML beta (95% CI)':fmt_ci(r.base_estimate,r.base_ci_lo,r.base_ci_hi),'Complete beta (95% CI)':fmt_ci(r.variant_estimate,r.variant_ci_lo,r.variant_ci_hi),'P complete':fmt_p(r.variant_p)})
        sheets['Word C complete case']=pd.DataFrame(compact)
        p=sheets['B weighted profiles'];p=p[p.universe=='All examined adults 60+']
        sheets['Word B profiles']=p.pivot(index=['variable','category'],columns='cognitive_availability',values='value').reset_index().round(3)
    with pd.ExcelWriter(path,engine='openpyxl') as w:
        for title,d in sheets.items():d.to_excel(w,sheet_name=title,index=False)
        pd.DataFrame({'Notes':notes}).to_excel(w,sheet_name='Notes',index=False)
        for ws in w.book:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24445E');c.alignment=Alignment(wrap_text=True)
            ws.row_dimensions[1].height=36
            for col in ws.columns:
                letter=col[0].column_letter;ws.column_dimensions[letter].width=min(46,max(13,max(len(str(c.value or '')) for c in list(col)[:100])+2))
            if ws.title=='Notes':
                ws.column_dimensions['A'].width=130
                for row in ws:
                    row[0].alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row[0].row].height=65
plan=json.loads((Q/'cycle_analysis_plan.json').read_text());retired=set(json.loads((Q/'cycle_contrast_plan.json').read_text())['retired_models']);cr=[];fr=[];all_est=[];all_eff=[]
for r in plan['models']:
    if r['model'] in retired:continue
    e,f,q,t=read('cycle_audit',r['model']);all_est.append(e);all_eff.append(f)
    if r['base'].startswith('p1_'):
        eb,fb,qb,tb=read('',r['base']);assert q['n']==qb['n']
        for variant,q0,t0 in [('base',qb,tb),('cycle adjusted',q,t)]:
            row={'model':r['base'],'variant':variant,'n':q0['n']}
            if 'Wald Test of Parameter Constraints' in t0:
                row['wald'],row['df'],row['p']=wald_result(t0)
            fr.append(row)
    else:
        for sol in ['unstandardized','STDYX'] if r['base'].startswith('m_') else ['unstandardized']:
            cr+=compare('cycle_audit',r['model'],r['base'],r['variant'],sol)
cycle=pd.DataFrame(cr);functional=pd.DataFrame(fr)
assert len([r for r in QC if r['folder']=='cycle_audit'])==54
pd.concat(all_est).to_csv(Q/'cycle_all_estimates.csv',index=False);pd.concat(all_eff).to_csv(Q/'cycle_all_effects.csv',index=False)
# Independent-cycle differences: common-scale biological paths only.
ct=[]
for (stem,quantity),g in cycle[(cycle.solution=='unstandardized')&cycle.variant.isin(['cycleG','cycleH'])&cycle.quantity.isin(['a1','a2','age_interaction'])].groupby(['model','quantity']):
    assert len(g)==2;gg=g.set_index('variant');a,b=gg.loc['cycleG'],gg.loc['cycleH'];delta=b.variant_estimate-a.variant_estimate;se=math.hypot(a.variant_se,b.variant_se)
    ct.append({'model':stem,'quantity':quantity,'G_b':a.variant_estimate,'H_b':b.variant_estimate,'H_minus_G':delta,'SE_difference':se,'CI_low':delta-1.96*se,'CI_high':delta+1.96*se,'P_exploratory':math.erfc(abs(delta/se)/math.sqrt(2))})
contrasts=pd.DataFrame(ct)
pooled=[]
for spec in json.loads((Q/'cycle_contrast_plan.json').read_text())['models']:
    e,f,q,t=read('cycle_contrasts',spec['model']);assert (q['strata'],q['psus'])==(29,61)
    u=e[e.solution=='unstandardized'];row=dict(spec,n=q['n'])
    for label,section,parameter in [('G',spec['outcome']+' ON',spec['exposure']),('H','New/Additional Parameters','BH'),('difference',spec['outcome']+' ON','C0')]:
        a=u[(u.section==section)&(u.parameter==parameter)];assert len(a)==1,(spec,a)
        for c in ['estimate','se','ci_lo','ci_hi','p']:row[label+'_'+c]=float(a.iloc[0][c])
    row['wald'],row['df'],row['wald_p']=wald_result(t);assert row['df']==1
    pooled.append(row)
pooled=pd.DataFrame(pooled);pooled.to_csv(Q/'cycle_pooled_contrasts.csv',index=False)
d=pd.read_csv(P/'data/processed/8_nhanes_analytic.csv',low_memory=False);d=d[d.RIDAGEYR>=20]
cy=[]
for cyc,g in d.groupby('CYCLE'):
    for c in ['LBXRDW','LBXWBCSI','KDM_BA','KDM_BA_advance','SII_log']:
        z=g.loc[g[c].notna()&(g.WTMEC4YR>0)];w=z.WTMEC4YR
        cy.append({'cycle':cyc,'variable':c,'n':len(z),'MEC_weighted_mean':np.average(z[c],weights=w),'unweighted_mean':z[c].mean(),'unweighted_SD':z[c].std()})
cycle.to_csv(Q/'cycle_comparisons.csv',index=False);contrasts.to_csv(Q/'cycle_contrasts.csv',index=False)
notes=[
 'Cycle sensitivity audit; candidate Supplementary Table S22. Baseline numbers are preserved. Original 2011–2014 definitions, scaling, weights, MLR and subpopulation design are retained.',
 'A: cycle-adjusted versus baseline. A binary period term enters every structural regression, including both biological mediators and the cognitive/depressive outcome. STDYX and common-file unstandardized results are separate; do not mix them. A one-SD latent outcome is not fixed across separate fits.',
 'B: functional-form tests compare the same spline/quintile specifications with and without the cycle term. The 3-df and 4-df Wald tests retain the original knots and hypotheses. These are a1-equation sensitivity analyses; they do not estimate nonlinear mediation.',
 'C: pooled single-outcome regressions allow the exposure, all covariate slopes and intercepts to differ by cycle. G and H slopes and their cycle interaction are on the common file scale; the formal test is a 1-df design-corrected Wald test. These diagnostic biological-measure associations use outcome-specific samples and are not full-SEM indirect paths. C2: cycle-specific age-interaction differences use SE_difference = sqrt(SE_G²+SE_H²) for independent cycles, printed-output precision, and exploratory normal tests.',
 'D: weighted means are descriptive, without hypothesis tests. Population 20+, observed value, positive four-year MEC weight. Unit RDW %, WBC 10^9/L, KDM years; log SII natural logarithm.',
 'E: cycle-specific age-interaction estimates. All 18 split-cycle full primary SEMs were retired from inference after eight showed design-rank/standard-error warnings; those runs remain in the audit archive. No cognitive or depressive full-SEM cycle contrast is reported. Replacement single-outcome models are fixed in cycle_contrast_plan.json.',
 'CDC documents a mid-2012 CBC instrument transition affecting RDW and WBC; both enter KDM. A cycle indicator cannot identify or repair the within-cycle transition, and period effects also combine secular and sample differences. No unsupported laboratory correction was applied.',
 'All P values are unadjusted sensitivity/diagnostic P values; 0.000 in numeric cells means Mplus printed <0.0005, not a literal zero. Coefficients and limits retain the precision of the output. Primary multiplicity handling remains unchanged.',
 'Source: https://wwwn.cdc.gov/Nchs/Data/Nhanes/Public/2013/DataFiles/CBC_H.htm; models and prospective audit scope: scripts/9_mplus_models/cycle_audit/plan.json.'
]
book(P/'result/supplementary/SuppTableS22_CycleSensitivity.xlsx',{'A cycle adjustment':cycle[cycle.variant=='adjusted'],'B functional form':functional,'C pooled cycle contrasts':pooled,'C2 age cycle contrasts':contrasts,'D CBC descriptive':pd.DataFrame(cy),'E cycle specific':cycle[cycle.variant!='adjusted'],'QC':pd.DataFrame([r for r in QC if r['folder'] in ['cycle_audit','cycle_contrasts']])},notes)
# Selection: all older examined adults, and those meeting RA-model exogenous criteria.
cols=re.findall(r'^#\s+\d+\s+(\w+)\s*$',(P/'data/mplus/nhanes_sem_p1_vars.txt').read_text(),re.M)
pd1=pd.read_csv(P/'data/mplus/nhanes_sem_p1.dat',sep=r'\s+',names=cols,na_values=-9999).set_index('SEQN')
o=d[d.RIDAGEYR>=60].copy();o=o[o.WTMEC4YR>0].set_index('SEQN')
o['cog_n']=o[['memory_score','fluency_score','speed_score']].notna().sum(axis=1)
o['exog_RA_ok']=pd1[['RA','AGE','SEX','BMI','EDU','WEAR_T','SLEEP_H']].notna().all(axis=1).reindex(o.index).fillna(False)
av=[];profile=[]
for universe,sub in [('All examined adults 60+',o),('RA model exogenous eligible',o[o.exog_RA_ok])]:
    for k,g in sub.groupby('cog_n'):
        av.append({'universe':universe,'observed_cognitive_scores':int(k),'n':len(g),'weighted_percent':100*g.WTMEC4YR.sum()/sub.WTMEC4YR.sum()})
    for k,g in [('none',sub[sub.cog_n==0]),('partial',sub[sub.cog_n.isin([1,2])]),('all three',sub[sub.cog_n==3])]:
        for v in ['RIDAGEYR','BMXBMI','DMDEDUC2','RA','IS','IV','KDM_BA_advance','SII_log']:
            z=g.loc[g[v].notna()];z=z[z[v].between(1,5)] if v=='DMDEDUC2' else z
            profile.append({'universe':universe,'cognitive_availability':k,'variable':v,'category':'weighted mean','n_observed':len(z),'value':np.average(z[v],weights=z.WTMEC4YR) if len(z) else np.nan})
        for v in ['RIAGENDR','DMDEDUC2','RIDRETH3','CYCLE']:
            z=g[g[v].notna()];z=z[z[v].between(1,5)] if v=='DMDEDUC2' else z
            for category,h in z.groupby(v):profile.append({'universe':universe,'cognitive_availability':k,'variable':v,'category':str(category)+' (%)','n_observed':len(h),'value':100*h.WTMEC4YR.sum()/z.WTMEC4YR.sum()})
sel=[]
for r in json.loads((Q/'selection_analysis_plan.json').read_text())['models']:
    e,f,q,t=read('selection_audit',r['model']);assert q['n']==r['n_complete'],(r,q)
    for sol in ['unstandardized','STDYX']:sel+=compare('selection_audit',r['model'],r['base'],'complete case',sol)
selection=pd.DataFrame(sel);selection.to_csv(Q/'selection_comparisons.csv',index=False)
notes2=[
 'Cognition availability and complete-case audit; candidate Supplementary Table S23. Availability is assessed among examined adults 60+ with positive MEC weights, with a second restricted universe meeting the RA-model exogenous-variable criteria. Zero test scores are valid when administration status supports them.',
 'A counts 0, 1, 2 or 3 observed scores (memory composite, animal fluency, digit-symbol speed); weighted percentages use all adults in the indicated universe. B gives descriptive weighted means or percentages among nonmissing values, separately by cognitive availability. No significance screening is used.',
 'C compares each of the three baseline cognition models with an otherwise identical model restricted to participants with all modeled exposure, covariates, mediators and cognitive indicators observed. All models use MLR and the original complex survey design and common input scales; the latent factor variance/orientation conventions remain unchanged.',
 'Availability differences indicate observed selection. This analysis does not establish missing at random, remove nonresponse bias, or prove that complete cases are unbiased. FIML can retain records lacking all cognitive indicators if a modeled mediator is observed; the reported SEM n is not necessarily the number contributing a cognitive score.',
 'Cognition testing was offered at different points in the MEC interview across periods; language/proxy/administration selection and top-coded age remain limitations. Period sensitivity and complete-case refits do not establish measurement invariance or repair unobserved outcomes.',
 'STDYX coefficients are sample-standardized within each fit. Unstandardized coefficients retain the common observed-variable scales but cognitive factor loadings/scales are reestimated in each sample. Compare direction, magnitude and uncertainty with these qualifications; all P values are unadjusted.',
 'Source: NHANES CFQ_G/CFQ_H local codebooks and scripts/9_mplus_models/selection_audit/plan.json. A numeric P of 0.000 denotes the output rounding threshold, not a literal zero.'
]
book(P/'result/supplementary/SuppTableS23_CognitiveSelection.xlsx',{'A availability':pd.DataFrame(av),'B weighted profiles':pd.DataFrame(profile),'C complete case':selection,'QC':pd.DataFrame([r for r in QC if r['folder']=='selection_audit'])},notes2)
pd.DataFrame(QC).to_csv(Q/'sensitivity_verified_outputs.csv',index=False)
summary={'status':'PASS','cycle_models':66,'retired_split_cycle_SEMs':18,'complete_case_models':3,'cycle_adjustment_max_abs_STDYX_change':float(cycle[(cycle.variant=='adjusted')&(cycle.solution=='STDYX')&(cycle.model.str.endswith(('_ra','_is','_iv')))].delta.abs().max()),'between_cycle_exploratory_P_below_005':pooled[pooled.wald_p<.05].to_dict('records'),'age_cycle_contrasts_P_below_005':contrasts[contrasts.P_exploratory<.05].to_dict('records'),'functional_tests':functional.dropna(subset=['p']).to_dict('records'),'complete_case_STDYX':selection[selection.solution=='STDYX'].to_dict('records'),'cognitive_availability':av}
(Q/'sensitivity_summary.json').write_text(json.dumps(summary,indent=2));print(json.dumps({k:v for k,v in summary.items() if k not in ['complete_case_STDYX','functional_tests']},indent=2))
