"""P1 report: design-corrected functional-form, age-interaction and
education-coding models -> Supplementary Table S21.

    python p1_report.py

Reads scripts/9_mplus_models/p1_*.out and m_*_edud.out (written by Mplus
from the inputs of 9_mplus_models/build_p1.py, run with mplus_run.py so that
each output carries a verified .run.json record) and writes
result/supplementary/SuppTableS21_P1_FunctionalForm.xlsx with four panels:

  Panel A  linear versus natural-spline a1 equation (RA, IV): the linear
           coefficient, the three non-linear spline coefficients and the
           design-corrected joint Wald test (MODEL TEST, 3 df)
  Panel B  spline x age: the joint Wald test of all four age-interaction
           terms (4 df), the test of the three non-linear interaction terms
           (3 df), and the linear x-by-age coefficient inside the spline model,
           next to the Table 3 interaction from the linear model
  Panel C  quintile indicators with common cut-points (Q2-Q5 against Q1) and
           the joint test of their age interactions (4 df)
  Panel D  the nine primary mediation models with education as four
           indicators instead of one 1-5 term: a1, b1, c', indirect via KDM-BA
           and total, STDYX, next to the primary single-term estimates

Every coefficient is on the scale of the model it comes from: Panels A-C are
unstandardised coefficients of the reference-standardised file (as Table 3);
Panel D is STDYX (as Table 2). Every output is validated by
mplus_output_qc before a number is read (run record, hashes, termination,
standard errors, input echo, Wald df, positive loadings); the QC sheet lists
the checks. This is the F1-freeze version of the report (2026-10-19).
"""
import re
import sys
import math
from functools import lru_cache
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
sys.path.insert(0, str(HERE))
from mplus_parse import parse_estimates, parse_effects  # noqa: E402
from mplus_export import fmt_ci, fmt_p  # noqa: E402
from mplus_output_qc import read_verified, validate_output, wald_result, sha256  # noqa: E402

MODELS = HERE / '9_mplus_models'
OUT = BASE / 'result' / 'supplementary' / 'SuppTableS21_P1_FunctionalForm.xlsx'
STEM = 'SuppTableS21_P1_FunctionalForm'
X = {'ra': 'RA', 'iv': 'IV'}
NAME = {'RA': 'Relative amplitude (RA)', 'IV': 'Intradaily variability (IV)'}


QC = {}


@lru_cache(maxsize=None)
def read(stem):
    f = MODELS / f'{stem}.out'
    if not f.exists():
        raise SystemExit(f"{f.name} is missing: run  python mplus_run.py 'p1_*.inp' && python mplus_run.py '*_edud.inp'  in scripts/9_mplus_models first")
    inp = f.with_suffix('.inp')
    if stem.startswith('p1_') or stem.endswith('_edud'):
        t, checked = read_verified(inp)
    else:
        t = f.read_text()
        checked = validate_output(inp, t)
    if (checked['strata'], checked['psus']) != (29, 61):
        raise ValueError(f'{stem}: unexpected survey design counts')
    QC[stem] = dict(model=stem, **checked, input_sha256=sha256(inp), output_sha256=sha256(f))
    return t


def wald(text):
    return wald_result(text)


def wald_p(p):
    """Retain the four-decimal precision of Mplus MODEL TEST output."""
    return '<0.001' if p < .001 else f'{p:.4f}'


def n_obs(text):
    m = re.search(r'Number of observations\s+(\d+)', text)
    return int(m.group(1)) if m else None


def coef(est, section, param, solution='unstandardized'):
    d = est[(est.solution == solution) & (est.section == section) & (est.parameter == param)]
    if len(d) != 1:
        raise ValueError(f'Expected exactly one {solution} {section} {param}; found {len(d)}')
    r = d.iloc[0]
    if not all(math.isfinite(float(v)) for v in (r.estimate, r.ci_lo, r.ci_hi, r.p)):
        raise ValueError(f'Nonfinite coefficient or interval: {section} {param}')
    return r.estimate, r.ci_lo, r.ci_hi, r.p


def ci(c):
    return fmt_ci(c[0], c[1], c[2]) if c else '—'


def panel_a():
    rows = []
    for k, x in X.items():
        lin = read(f'p1_{k}_linear'); spl = read(f'p1_{k}_spline')
        if n_obs(lin) != n_obs(spl):
            raise ValueError(f'{x}: linear and spline samples differ')
        el = pd.DataFrame(parse_estimates(lin, f'p1_{k}_linear')); es = pd.DataFrame(parse_estimates(spl, f'p1_{k}_spline'))
        w, df, p = wald(spl)
        b1l = coef(el, 'BA_ADV ON', x); b1s = coef(es, 'BA_ADV ON', x)
        rows.append({'Rhythm metric': NAME[x], 'n': f'{n_obs(spl):,}',
                     'Linear model: b (95% CI)': ci(b1l), 'Linear model: P': fmt_p(b1l[3]) if b1l else '—',
                     'Spline model: linear term b (95% CI)': ci(b1s),
                     'Spline model: N3 (95% CI)': ci(coef(es, 'BA_ADV ON', f'{x}_N3')),
                     'Spline model: N4 (95% CI)': ci(coef(es, 'BA_ADV ON', f'{x}_N4')),
                     'Spline model: N5 (95% CI)': ci(coef(es, 'BA_ADV ON', f'{x}_N5')),
                     'Wald χ² (non-linear terms = 0)': f'{w:.3f}' if w is not None else '—', 'df': df if df else '—',
                     'P': wald_p(p) if p is not None else '—'})
    return pd.DataFrame(rows)


def panel_b():
    rows = []
    t3 = pd.read_excel(BASE / 'result/main/Table3_Age_Moderation.xlsx', 'Table 3')
    t3['Exposure'] = t3['Exposure'].ffill()
    bcol = 'b (95% CI)' if 'b (95% CI)' in t3.columns else 'β (95% CI)'
    for k, x in X.items():
        full = read(f'p1_{k}_spline_age'); nl = read(f'p1_{k}_spline_age_nl')
        if n_obs(full) != n_obs(nl):
            raise ValueError(f'{x}: the two spline-interaction fits have different samples')
        e = pd.DataFrame(parse_estimates(full, f'p1_{k}_spline_age'))
        w4, df4, p4 = wald(full); w3, df3, p3 = wald(nl)
        i1 = coef(e, 'BA_ADV ON', f'{x}_AGE')
        lin = t3[(t3.Exposure == x) & t3.Term.astype(str).str.contains('× age')]
        rows.append({'Rhythm metric': NAME[x], 'n': f'{n_obs(full):,}',
                     'Linear model (Table 3): x × age b (95% CI)': lin.iloc[0][bcol] if len(lin) else '—',
                     'Spline model: x × age b (95% CI)': ci(i1), 'Spline model: x × age P': fmt_p(i1[3]) if i1 else '—',
                     'Spline model: N3 × age (95% CI)': ci(coef(e, 'BA_ADV ON', f'{x}N3AGE')),
                     'Spline model: N4 × age (95% CI)': ci(coef(e, 'BA_ADV ON', f'{x}N4AGE')),
                     'Spline model: N5 × age (95% CI)': ci(coef(e, 'BA_ADV ON', f'{x}N5AGE')),
                     'Wald χ² (all four age terms = 0)': f'{w4:.3f}' if w4 is not None else '—', 'df (all)': df4 or '—',
                     'P (all)': wald_p(p4) if p4 is not None else '—',
                     'Wald χ² (three non-linear age terms = 0)': f'{w3:.3f}' if w3 is not None else '—', 'df (non-linear)': df3 or '—',
                     'P (non-linear)': wald_p(p3) if p3 is not None else '—'})
    return pd.DataFrame(rows)


def panel_c():
    rows = []
    for k, x in X.items():
        q = read(f'p1_{k}_quint'); qa = read(f'p1_{k}_quint_age')
        if n_obs(q) != n_obs(qa):
            raise ValueError(f'{x}: quintile and quintile-interaction samples differ')
        eq = pd.DataFrame(parse_estimates(q, f'p1_{k}_quint')); ea = pd.DataFrame(parse_estimates(qa, f'p1_{k}_quint_age'))
        w, df, p = wald(qa)
        for j in (2, 3, 4, 5):
            c = coef(eq, 'BA_ADV ON', f'{x}_Q{j}'); ca = coef(ea, 'BA_ADV ON', f'{x}Q{j}AGE')
            rows.append({'Rhythm metric': NAME[x], 'n': f'{n_obs(q):,}', 'Contrast': f'Q{j} vs Q1',
                         'b (95% CI), no interaction': ci(c), 'P': fmt_p(c[3]) if c else '—',
                         'Q × age b (95% CI)': ci(ca), 'Q × age P': fmt_p(ca[3]) if ca else '—',
                         'Wald χ² (four Q × age terms = 0)': f'{w:.3f}' if (w is not None and j == 2) else '',
                         'df': df if (df and j == 2) else '', 'P (joint)': wald_p(p) if (p is not None and j == 2) else ''})
    return pd.DataFrame(rows)


def panel_d():
    rows = []
    cells = {('dep', 'young'): ('PHQ-9', '20–59'), ('dep', 'old'): ('PHQ-9', '≥60'), ('cog', 'old'): ('cognition', '≥60')}
    for exp in ('RA', 'IS', 'IV'):
        for (outcome, cell), (ylab, agelab) in cells.items():
            stem = f'm_{outcome}_{cell}_{exp.lower()}'
            prim = read(stem); edud = read(f'{stem}_edud')
            if n_obs(prim) != n_obs(edud):
                raise ValueError(f'{stem}: education recoding changed the sample')
            ep = pd.DataFrame(parse_estimates(prim, stem)); ed = pd.DataFrame(parse_estimates(edud, f'{stem}_edud'))
            fp = pd.DataFrame(parse_effects(prim, stem)); fd = pd.DataFrame(parse_effects(edud, f'{stem}_edud'))
            y = 'COG' if outcome == 'cog' else 'PHQ9'
            def eff(f, kind, via=None):
                d = f[(f.solution == 'STDYX') & (f.effect == kind)]
                d = d[d.via == via] if via else d[d.via.isna()]
                if len(d) != 1:
                    raise ValueError(f'{stem}: missing or duplicate {kind} {via}')
                vals = (d.iloc[0].estimate, d.iloc[0].ci_lo, d.iloc[0].ci_hi, d.iloc[0].p)
                if not all(math.isfinite(float(v)) for v in vals):
                    raise ValueError(f'{stem}: nonfinite {kind} {via}')
                return vals
            for label, getter in [('a1: rhythm → KDM-BA advancement', lambda e, f: coef(e, 'BA_ADV ON', exp, 'STDYX')),
                                  ('b1: KDM-BA advancement → outcome', lambda e, f: coef(e, f'{y} ON', 'BA_ADV', 'STDYX')),
                                  ("c′: direct", lambda e, f: eff(f, 'direct')),
                                  ('indirect via KDM-BA advancement', lambda e, f: eff(f, 'indirect', 'BA_ADV')),
                                  ('total', lambda e, f: eff(f, 'total'))]:
                cp, cd = getter(ep, fp), getter(ed, fd)
                rows.append({'Outcome': ylab, 'Age group': agelab, 'Rhythm metric': exp, 'Quantity (STDYX)': label,
                             'n (single term)': f'{n_obs(prim):,}', 'Education as single 1–5 term: β (95% CI)': ci(cp), 'P': fmt_p(cp[3]) if cp else '—',
                             'n (four indicators)': f'{n_obs(edud):,}', 'Education as four indicators: β (95% CI)': ci(cd), 'P ': fmt_p(cd[3]) if cd else '—',
                             'Δβ': f'{cd[0] - cp[0]:+.3f}' if (cp and cd) else '—'})
    return pd.DataFrame(rows)


NOTES = [
    'Supplementary Table S21 | Design-corrected functional-form, age-interaction and education-coding analyses (P1).',
    'All models: Mplus 9, TYPE = COMPLEX, MLR, the four-year MEC weight, 29 strata and 61 PSUs, SUBPOPULATION = adults aged 20 and over; '
    'dependent variable KDM biological-age advancement (the a1 leg of the structural models) unless stated; covariates age, sex, body mass index, '
    'education, accelerometer wear time and sleep duration. MODEL TEST gives the design-corrected joint Wald test of the listed parameter constraints.',
    'Panels A-C: coefficients are from the unstandardised solution, not STDYX. The linear exposure, age, and biological-age advancement retain '
    'the common reference scaling used in Table 3 (n = 7,882); sex and quintile indicators retain their original coding. The natural-spline '
    'basis is the one used in Supplementary Table S10 (Hastie et al., ESL eq. 5.4-5.5) in the raw metric, with boundary knots at the 5th and 95th '
    'percentiles and three interior knots at the quartiles, computed once on adults aged 20+ with the metric observed and applied to every model; '
    'the three non-linear basis functions are z-scored on that sample (an affine change that leaves the tests unchanged). The linear term is the '
    'z-scored metric of the file; its coefficient in the spline model is not the slope of the metric, because the non-linear terms are correlated with it.',
    'Panel A tests whether the three non-linear terms add to the straight line (3 df). Panel B adds age and the four products of age with the spline '
    'terms: the 4-df test asks whether age modifies the association in any way (shape or slope); the 3-df test asks whether age modifies its shape; '
    'the x × age coefficient is the age interaction of the linear part inside the spline model and is compared with the Table 3 interaction from '
    'the model without spline terms. Panel C uses quintile indicators of the raw metric with cut-points common to every model (reference: first '
    'quintile); the 4-df test is on the four quintile × age products.',
    'Panel D refits the nine primary mediation models with education entered as four indicators (NHANES categories 2-5 against less than 9th grade) '
    'instead of the single standardised 1-5 term; STDYX coefficients as in Table 2 and Supplementary Tables S3-S4. Participants with missing '
    'education are excluded in both specifications; identical n is checked before export. Panel D P values are unadjusted; they are not the '
    'within-model Holm-adjusted indirect-association P values used in the primary report.',
    'These analyses were added during manuscript revision in response to the supportive diagnostics of Supplementary Table S10 and are reported '
    'as sensitivity analyses. Model files: p1_{ra,iv}_{linear,spline,spline_age,spline_age_nl,quint,quint_age}.inp and m_*_edud.inp.',
    'KDM, Klemera-Doubal method; PHQ-9, 9-item Patient Health Questionnaire; RA, relative amplitude; IV, intradaily variability; IS, interdaily stability.',
]


def main():
    A, B, C, D = panel_a(), panel_b(), panel_c(), panel_d()
    if len(QC) != 30:
        raise ValueError(f'Expected 21 P1 and nine primary comparison outputs; read {len(QC)}')
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
        A.to_excel(xw, sheet_name='S21 Panel A linear vs spline', index=False)
        B.to_excel(xw, sheet_name='S21 Panel B spline x age', index=False)
        C.to_excel(xw, sheet_name='S21 Panel C quintile x age', index=False)
        D.to_excel(xw, sheet_name='S21 Panel D education coding', index=False)
        pd.DataFrame({'Notes': NOTES}).to_excel(xw, sheet_name='Notes', index=False)
        pd.DataFrame(QC.values()).to_excel(xw, sheet_name='QC verified outputs', index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = 'C2' if ws.title.startswith('S21') else 'A2'
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 68 if ws.title.startswith('S21') else 28
            for cell in ws[1]:
                cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='24445E')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                h = str(cell.value)
                width = (12 if h.startswith(('P', 'df', 'n', 'Δβ')) and 'coefficient' not in h
                         else 34 if 'sha256' in h or h in ('Notes', 'Quantity (STDYX)', 'Rhythm metric') else 28)
                ws.column_dimensions[cell.column_letter].width = 100 if h == 'Notes' else width
            for row in ws.iter_rows(min_row=2):
                ws.row_dimensions[row[0].row].height = 36 if ws.title.startswith('S21') else 30
                for cell in row:
                    cell.font = Font(name='Arial', size=10)
                    cell.alignment = Alignment(wrap_text=True, vertical='center')
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill('solid', fgColor='F1F5F8')
            if ws.title == 'Notes':
                for row in ws.iter_rows(min_row=2):
                    ws.row_dimensions[row[0].row].height = max(32, 15 * (len(str(row[0].value)) // 115 + 2))
    pd.set_option('display.width', 250)
    print(A.to_string(index=False)); print(); print(B.to_string(index=False)); print(); print(C.to_string(index=False))
    print(f'\nwrote {OUT.relative_to(BASE)}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
